#!/usr/bin/env python3
"""Build the UFI Phone diploma project as a polished DOCX artifact."""

from __future__ import annotations

import importlib.util
import re
from datetime import datetime
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
CONTENT_PATH = ROOT / "diploma" / "diploma_content.md"
OUTPUT_DIR = ROOT / "artifacts" / "diploma"
WORK_DIR = OUTPUT_DIR / "_work"
OUTPUT_PATH = OUTPUT_DIR / "Дипломный проект UFI Phone Куватов Руслан Бахтиярович.docx"
ASSET_DIR = ROOT / "presentations" / "assets"
REFERENCE_BUILDER = Path(
    "/home/Hollow/Documents/Codex/2026-09-06/i/recovered-projects/"
    "Codex/2026-07-03/sure-for-a-perfect-online-ielts/ielts-lms-platform/"
    "scripts/build_diploma.py"
)

spec = importlib.util.spec_from_file_location("diploma_reference_builder", REFERENCE_BUILDER)
if spec is None or spec.loader is None:
    raise RuntimeError(f"Cannot load reference builder: {REFERENCE_BUILDER}")
base = importlib.util.module_from_spec(spec)
spec.loader.exec_module(base)

base.ROOT = ROOT
base.CONTENT_PATH = CONTENT_PATH
base.OUTPUT_DIR = OUTPUT_DIR
base.WORK_DIR = WORK_DIR
base.OUTPUT_PATH = OUTPUT_PATH


TABLES = {
    "alternatives": {
        "caption": "Сравнение способов использования SIM-карты в модеме",
        "headers": ["Подход", "Вызовы и звук", "SMS", "Ограничение"],
        "rows": [
            ["Штатный web-интерфейс UFI", "Обычно отсутствуют", "Зависит от прошивки", "Ориентирован на доступ в Интернет"],
            ["AT-команды с компьютера", "Управление возможно частично", "Чтение и отправка", "Нужны драйверы, USB-доступ и технический интерфейс"],
            ["SMS → Telegram", "Не поддерживаются", "Облачная пересылка", "Содержимое покидает локальную сеть"],
            ["Обычный смартфон", "Полная телефония", "Полная поддержка", "SIM нельзя одновременно использовать в модеме"],
            ["UFI Phone", "Входящие/исходящие вызовы и двусторонний звук", "Локальные диалоги и отправка", "Требуется точный аппаратный профиль и ADB-установка"],
        ],
        "widths": [3.0, 4.5, 3.2, 4.8],
    },
    "hardware": {
        "caption": "Параметры проверенного аппаратного профиля",
        "headers": ["Параметр", "Значение", "Назначение проверки"],
        "rows": [
            ["Модель", "UFI003 / LTE 4G Wi-Fi Dongle", "Идентификация семейства устройства"],
            ["USB ID", "05c6:90b4", "Отсечение внешне похожих модемов"],
            ["SoC", "Qualcomm MSM8916, ARMv7", "Совместимость двоичного окружения"],
            ["ОС модема", "Android 4.4.4, API 19", "Доступные системные API"],
            ["Baseband", "UFI003_CT 20220903", "Семейство радиопрошивки"],
            ["LAN", "192.168.100.1/24", "Изолированная связь клиентов с шлюзом"],
            ["Голос", "2G/3G circuit-switched fallback", "Честная граница: IMS/VoLTE отсутствует"],
        ],
        "widths": [3.2, 5.0, 7.3],
    },
    "operators": {
        "caption": "Статус совместимости с операторами Узбекистана",
        "headers": ["Оператор", "Статус UFI003", "Решение"],
        "rows": [
            ["Ucell", "Проверены LTE-данные, SMS, вызовы и двусторонний звук", "Референсный поддерживаемый оператор"],
            ["Mobiuz", "Архитектура fallback совместима, SIM-тест не завершён", "Полная приёмочная матрица до заявления поддержки"],
            ["Uzmobile", "Ожидается совместимость, проектный SIM-тест не выполнен", "Проверить вызовы и возврат данных в LTE"],
            ["Humans", "MVNO на сети Uzmobile; профиль SIM может отличаться", "Проверять отдельно"],
            ["Beeline", "IMS модема отсутствует; CS-поведение не подтверждено", "Тестировать тариф и fallback"],
            ["Perfectum", "CDMA/5G SA/VoNR-пути не соответствуют UFI003", "Не поддерживать этим аппаратным профилем"],
        ],
        "widths": [2.6, 6.4, 6.5],
    },
    "requirements": {
        "caption": "Основные функциональные требования",
        "headers": ["Код", "Требование", "Критерий приёмки"],
        "rows": [
            ["FR-01", "Показывать состояние модема и готовность к вызову", "Клиент различает mode 9, LTE-only и отсутствие связи"],
            ["FR-02", "Принимать и инициировать обычные сотовые вызовы", "Отображаются номер, направление и состояние; доступны Answer/Hang up"],
            ["FR-03", "Передавать двусторонний звук", "Разговор слышен в обоих направлениях; занят второй аудиоклиент отклоняется"],
            ["FR-04", "Синхронизировать и отправлять SMS", "До 250 сообщений доступны по диалогам, Unicode сохраняется"],
            ["FR-05", "Сохранять локальную историю", "До 100 вызовов с временем, результатом и длительностью"],
            ["FR-06", "Автоматически запускаться после перезагрузки", "Шлюз, guard и монитор восстанавливают работу"],
            ["FR-07", "Безопасно устанавливать компоненты", "Неизвестный профиль отвергается до привилегированной установки"],
            ["FR-08", "Работать без обязательного облачного сервиса", "Основные вызовы и SMS остаются в частной LAN"],
        ],
        "widths": [1.6, 6.0, 7.9],
    },
    "quality": {
        "caption": "Нефункциональные требования и способы проверки",
        "headers": ["Код", "Характеристика", "Проектное решение"],
        "rows": [
            ["NFR-01", "Безопасность", "HMAC-SHA-256 challenge-response, случайный 256-битный token, привязка к частной /24"],
            ["NFR-02", "Конфиденциальность", "Нет аналитики и облачного relay; история хранится в приватных каталогах клиентов"],
            ["NFR-03", "Надёжность", "START_STICKY, boot receivers, сетевой guard и повтор приведения режима"],
            ["NFR-04", "Совместимость", "Точное совпадение product, SDK, baseband, USB ID и LAN-параметров"],
            ["NFR-05", "Сопровождаемость", "Разделение gateway, guard, клиентов, installer и аппаратных профилей"],
            ["NFR-06", "Переносимость клиента", "Android 8+, Linux и Windows используют общий протокол"],
            ["NFR-07", "Проверяемость", "Модульные/интеграционные тесты и CI guard версии, секретов и Action pinning"],
        ],
        "widths": [1.6, 3.7, 10.2],
    },
    "components": {
        "caption": "Компоненты программно-аппаратного комплекса",
        "headers": ["Компонент", "Среда", "Ответственность"],
        "rows": [
            ["Network Guard", "Android 4.4, uid phone", "Восстановление безопасного preferred network mode 9"],
            ["Voice Gateway", "Android 4.4, uid system", "Вызовы, SMS, состояние и три LAN-сокета"],
            ["UFI Phone Android", "Android 8+", "Calls, Keypad, Messages, уведомления и звук"],
            ["Desktop client", "Linux/Windows", "Полный интерфейс, история и локальный аудиобэкенд"],
            ["ufi_setup.py", "Linux/Windows с ADB", "Doctor, exact-match профиль, сборка и парное развёртывание"],
            ["ufi_voice.py", "Python", "Настройка, status, CLI и запуск desktop UI"],
            ["ufi_sms.py", "Linux/PyUSB", "Резервная работа с SMS через USB AT-интерфейс"],
            ["Telegram forwarder", "Android 4.4, опционально", "Совместимый облачный путь с устойчивой очередью"],
        ],
        "widths": [3.3, 4.0, 8.2],
    },
    "protocol": {
        "caption": "Контракт локального протокола UFI Phone",
        "headers": ["Канал", "Порт/формат", "Операции и ограничения"],
        "rows": [
            ["Control", "TCP 8765, UTF-8 lines", "PING, STATUS, DIAL, ANSWER, HANGUP, SMS_LIST, SMS_SEND, SMS_READ"],
            ["Downlink", "TCP 8766, PCM S16LE 8 kHz mono", "Звук сети к клиенту; один активный поток"],
            ["Uplink", "TCP 8767, PCM S16LE 48 kHz mono", "Микрофон клиента к модему; один активный поток"],
            ["Authentication", "HELLO 2 + 32-byte nonce", "HMAC-SHA-256 связывает token, port, nonce и command/stream"],
            ["Network scope", "192.168.100.0/24", "Шлюз отвергает клиентов вне частной LAN"],
            ["Limits", "12 sessions, bounded lines", "Защита от неограниченных соединений и payload"],
        ],
        "widths": [3.0, 5.2, 7.3],
    },
    "threats": {
        "caption": "Модель угроз и меры управления",
        "headers": ["Угроза", "Последствие", "Мера"],
        "rows": [
            ["Команда от постороннего клиента", "Несанкционированный вызов или чтение SMS", "Private /24, nonce и HMAC для каждой сессии"],
            ["Повтор перехваченной команды", "Replay управления", "Свежий 32-байтный nonce и command binding"],
            ["Утечка pairing token", "Доступ внутри LAN", "Генерация 256 бит, права пользователя, передача только через ADB"],
            ["Установка на похожий модем", "Bootloop или компрометация привилегий", "Exact-match профиль и проверка сертификата"],
            ["Набор экстренного номера", "Опасное ложное ожидание поддержки", "Emergency/service dialing блокируется"],
            ["Два аудиоклиента", "Смешение приватного разговора", "Atomic busy lock на uplink/downlink"],
            ["Облачная утечка SMS", "Раскрытие содержания", "Telegram выключен по умолчанию; основной путь локальный"],
        ],
        "widths": [4.0, 4.7, 6.8],
    },
    "tests": {
        "caption": "Результаты воспроизводимого тестового прогона",
        "headers": ["Область", "Проверка", "Результат 25.09.2026"],
        "rows": [
            ["Android manifests", "Согласованность standalone/Gradle capabilities", "Пройдено"],
            ["Desktop model", "Диалоги, непрочитанные, история и нормализация номера", "3 теста пройдены"],
            ["Hardware profiles", "Точное совпадение и отказ near-match", "3 теста пройдены"],
            ["SMS durability", "Границы polling, rotation и чтение архива", "2 теста пройдены"],
            ["Voice protocol", "Challenge-response, payload, адрес и отказ plaintext", "3 теста пройдены"],
            ["Итого unittest", "uv + pyusb 1.3.1", "12 из 12 пройдены за 1,022 с"],
            ["Project guard", "Версии, секреты, размеры и pins GitHub Actions", "Пройдено для версии 0.5.0"],
        ],
        "widths": [3.4, 7.2, 4.9],
    },
    "acceptance": {
        "caption": "Матрица приёмки нового модема или оператора",
        "headers": ["Этап", "Проверяемый результат", "Условие допуска"],
        "rows": [
            ["Идентификация", "USB ID, product, SDK, baseband, LAN", "Нет необъяснённого расхождения с профилем"],
            ["SMS", "Входящие, исходящие, multipart, Unicode, read state", "Повторяемо после reboot"],
            ["Вызовы", "Ring, dial, answer, hangup, missed", "Состояния и история согласованы"],
            ["Аудио", "Речь в обоих направлениях", "Нет второго одновременного клиента"],
            ["Радиорежим", "10 cold boots и возврат LTE после вызова", "Mode 9 восстанавливается автоматически"],
            ["Безопасность", "Wrong token и внешний адрес", "Доступ отклонён без раскрытия данных"],
            ["Оператор", "SIM/тариф, charging, APN и fallback", "Результат опубликован с firmware/tariff"],
        ],
        "widths": [3.0, 7.0, 5.5],
    },
    "safety": {
        "caption": "Опасные факторы и профилактические меры",
        "headers": ["Фактор", "Проявление", "Профилактика"],
        "rows": [
            ["Электрический", "Повреждённый USB-кабель или блок питания", "Исправный сертифицированный адаптер, осмотр, запрет ремонта под напряжением"],
            ["Тепловой", "Перегрев модема при длительной передаче", "Свободная вентиляция, негорючая поверхность, контроль температуры"],
            ["Пожарный", "Короткое замыкание или перегрузка", "Исправная защита, свободный выход, подходящий огнетушитель"],
            ["Эргономический", "Статическая поза и зрительная нагрузка", "Настройка рабочего места и регулярные перерывы"],
            ["Акустический", "Эхо и чрезмерная громкость", "Гарнитура, умеренный уровень, echo cancellation"],
            ["Информационный", "Раскрытие номера, SMS или разговора", "Private LAN, token, локальное хранение и минимизация журналов"],
            ["Радиочастотный", "Длительное размещение у тела", "Следовать документации устройства и обеспечивать дистанцию/вентиляцию"],
        ],
        "widths": [3.0, 6.1, 6.4],
    },
}

base.TABLES = TABLES


TOC_ENTRIES = [
    (0, "ВВЕДЕНИЕ", 6),
    (0, "ГЛАВА I. АНАЛИЗ ПРЕДМЕТНОЙ ОБЛАСТИ И ПОСТАНОВКА ЗАДАЧИ", 10),
    (1, "1.1 Назначение LTE UFI-модемов и проблема повторного использования SIM", 10),
    (1, "1.2 Аппаратно-программные особенности UFI003", 12),
    (1, "1.3 Сотовые вызовы, SMS и механизм fallback", 14),
    (1, "1.4 Анализ существующих решений", 16),
    (1, "1.5 Рынок и совместимость операторов Узбекистана", 17),
    (1, "1.6 Постановка задачи и требования", 19),
    (0, "ГЛАВА II. ПРОЕКТИРОВАНИЕ КОМПЛЕКСА UFI PHONE", 22),
    (1, "2.1 Пользователи, сценарии и границы системы", 22),
    (1, "2.2 Компонентная архитектура", 23),
    (1, "2.3 Проектирование аппаратных профилей и установки", 26),
    (1, "2.4 Локальный протокол и модель состояний", 28),
    (1, "2.5 Проектирование голосового и аудиоканала", 30),
    (1, "2.6 Проектирование SMS и пользовательских клиентов", 32),
    (1, "2.7 Безопасность, приватность и надёжность", 34),
    (0, "ГЛАВА III. РЕАЛИЗАЦИЯ И ТЕСТИРОВАНИЕ", 36),
    (1, "3.1 Реализация doctor, профилей и безопасной установки", 36),
    (1, "3.2 Реализация gateway и network guard", 37),
    (1, "3.3 Реализация вызовов и двустороннего звука", 38),
    (1, "3.4 Реализация SMS и клиентских интерфейсов", 39),
    (1, "3.5 Desktop, упаковка и поставка", 40),
    (1, "3.6 Тестирование и результаты", 42),
    (1, "3.7 Ограничения, приёмка и развитие", 44),
    (0, "ГЛАВА IV. БЕЗОПАСНОСТЬ ЖИЗНЕДЕЯТЕЛЬНОСТИ", 47),
    (1, "4.1 Анализ условий труда и рисков", 47),
    (1, "4.2 Эргономика и расчёт освещения", 48),
    (1, "4.3 Электрическая, тепловая и пожарная безопасность", 50),
    (1, "4.4 Радиочастотная и информационная безопасность", 51),
    (0, "ЗАКЛЮЧЕНИЕ", 53),
    (0, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 55),
]


FIGURE_CAPTIONS = {
    "context": "Контекст программно-аппаратного комплекса UFI Phone",
    "radio": "Переход радиосети при обычном сотовом вызове",
    "architecture": "Логические уровни архитектуры комплекса",
    "trust": "Граница доверия и сетевые каналы протокола",
    "install": "Алгоритм безопасной установки по аппаратному профилю",
    "call_sequence": "Последовательность обработки входящего вызова",
    "audio": "Двусторонний аудиотракт между модемом и клиентом",
    "sms": "Поток получения, хранения и отображения SMS",
    "ui": "Интерфейсы вызовов и сообщений UFI Phone",
    "deployment": "Варианты поставки компонентов комплекса",
    "quality": "Контуры проверки версии 0.5.0",
}


def add_title_page(doc: Document) -> None:
    base.add_centered(doc, "МИНИСТЕРСТВО ЦИФРОВЫХ ТЕХНОЛОГИЙ", size=14, bold=True)
    base.add_centered(doc, "РЕСПУБЛИКИ УЗБЕКИСТАН", size=14, bold=True, space_after=6)
    base.add_centered(doc, "ТАШКЕНТСКИЙ УНИВЕРСИТЕТ ИНФОРМАЦИОННЫХ", size=14, bold=True)
    base.add_centered(doc, "ТЕХНОЛОГИЙ ИМЕНИ МУХАММАДА АЛ-ХОРАЗМИЙ", size=14, bold=True, space_after=12)
    base.add_centered(doc, "ФАКУЛЬТЕТ КОМПЬЮТЕРНЫЙ ИНЖИНИРИНГ", size=14, bold=True, space_after=22)

    approval = doc.add_table(rows=5, cols=2)
    approval.alignment = WD_TABLE_ALIGNMENT.RIGHT
    approval.autofit = False
    lines = [
        "Допустить к защите",
        "Заведующий кафедрой",
        "«Компьютерные системы»",
        "____________ Рахимов М.Ф.",
        "«___» ______________ 2026 г.",
    ]
    for row in approval.rows:
        base.set_cell_width(row.cells[0], 9.7)
        base.set_cell_width(row.cells[1], 6.8)
    for index, text in enumerate(lines):
        p = approval.cell(index, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        base.set_run_font(r, size=14, bold=index == 0)
    base.remove_table_borders(approval)

    base.add_centered(doc, "ДИПЛОМНЫЙ ПРОЕКТ", size=20, bold=True, space_before=24, space_after=18)
    topic = doc.add_paragraph()
    topic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    topic.paragraph_format.first_line_indent = Cm(0)
    topic.paragraph_format.line_spacing = 1.5
    topic.paragraph_format.space_after = Pt(34)
    r = topic.add_run("Тема: ")
    base.set_run_font(r, size=14, bold=True)
    r = topic.add_run(
        "Разработка программно-аппаратного комплекса UFI Phone для управления "
        "голосовыми вызовами и SMS через LTE USB-модем"
    )
    base.set_run_font(r, size=14)

    people = doc.add_table(rows=8, cols=4)
    people.alignment = WD_TABLE_ALIGNMENT.CENTER
    people.autofit = False
    values = [
        ("", "Выпускник", "_______________", "Куватов Р. Б."),
        ("", "", "(подпись)", "(ФИО)"),
        ("", "Руководитель", "_______________", "Азамова С. Ф."),
        ("", "", "(подпись)", "(ФИО)"),
        ("", "Рецензент", "_______________", "______________"),
        ("", "", "(подпись)", "(ФИО)"),
        ("", "Консультант\nпо БЖД", "_______________", "______________"),
        ("", "", "(подпись)", "(ФИО)"),
    ]
    widths = [3.4, 3.5, 4.2, 5.4]
    for i, row in enumerate(people.rows):
        for j, cell in enumerate(row.cells):
            base.set_cell_width(cell, widths[j])
            base.set_cell_margins(cell, 15, 30, 15, 30)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(values[i][j])
            base.set_run_font(r, size=10 if i % 2 else 14)
            if i % 2 == 0 and j == 3 and "_" not in values[i][j]:
                r.underline = True
    base.remove_table_borders(people)
    base.add_centered(doc, "Ташкент – 2026", size=14, bold=True, space_before=36)
    doc.add_page_break()


def add_assignment_page(doc: Document) -> None:
    base.add_centered(doc, "МИНИСТЕРСТВО ЦИФРОВЫХ ТЕХНОЛОГИЙ", size=14, bold=True)
    base.add_centered(doc, "РЕСПУБЛИКИ УЗБЕКИСТАН", size=14, bold=True)
    base.add_centered(doc, "ТАШКЕНТСКИЙ УНИВЕРСИТЕТ ИНФОРМАЦИОННЫХ", size=14, bold=True)
    base.add_centered(doc, "ТЕХНОЛОГИЙ ИМЕНИ МУХАММАДА АЛ-ХОРАЗМИЙ", size=14, bold=True, space_after=8)
    base.add_labeled_line(doc, "Факультет:  ", "Компьютерный инжиниринг", italic=True)
    base.add_labeled_line(doc, "Кафедра:     ", "Компьютерные системы", italic=True)
    base.add_labeled_line(
        doc,
        "Направление (специальность):  ",
        "60610500 – Компьютерный инжиниринг («Компьютерный инжиниринг»)",
        space_after=8,
    )

    approval = doc.add_table(rows=5, cols=2)
    approval.alignment = WD_TABLE_ALIGNMENT.RIGHT
    approval.autofit = False
    lines = [
        "УТВЕРЖДАЮ",
        "Заведующий кафедрой",
        "«Компьютерные системы»",
        "____________ Рахимов М.Ф.",
        "«___» ______________ 2026 г.",
    ]
    for row in approval.rows:
        base.set_cell_width(row.cells[0], 9.5)
        base.set_cell_width(row.cells[1], 7.0)
    for index, text in enumerate(lines):
        p = approval.cell(index, 1).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        base.set_run_font(r, size=14, bold=index == 0)
    base.remove_table_borders(approval)

    base.add_centered(doc, "ЗАДАНИЕ", size=16, bold=True, space_before=14, space_after=6)
    base.add_centered(doc, "на дипломный проект", size=14, space_after=5)
    student = base.add_centered(doc, "Куватов Руслан Бахтиярович", size=14, bold=True)
    student.runs[0].underline = True
    base.add_centered(doc, "(фамилия, имя, отчество)", size=9, space_after=7)

    items = [
        ("1. Тема ДП: ", "Разработка программно-аппаратного комплекса UFI Phone для управления голосовыми вызовами и SMS через LTE USB-модем."),
        ("2. Утверждена приказом по университету от: ", "«___» __________ 2026 г. № ______."),
        ("3. Срок сдачи законченной работы: ", "______________________________."),
        ("4. Исходные данные к работе: ", "Исходный код и документация UFI Phone, проверенный модем UFI003, результаты испытаний, официальные материалы Android, операторов связи и нормативные источники."),
        ("5. Содержание расчётно-пояснительной записки: ", "Аннотация; введение; анализ предметной области; проектирование комплекса; реализация и тестирование; безопасность жизнедеятельности; заключение; список источников."),
        ("6. Перечень графического материала: ", "Архитектурные и сетевые схемы, алгоритмы, таблицы испытаний, интерфейсы приложения и презентация Microsoft PowerPoint."),
        ("7. Дата выдачи задания: ", "«___» ______________ 2026 г."),
    ]
    for index, (label, value) in enumerate(items):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(label)
        base.set_run_font(r, size=12.5, bold=True)
        r = p.add_run(value)
        base.set_run_font(r, size=12.5, italic=index in (3, 4, 5))
        if index > 0:
            r.underline = True

    sign = doc.add_table(rows=2, cols=3)
    sign.alignment = WD_TABLE_ALIGNMENT.RIGHT
    sign.autofit = False
    sign_values = [("", "Руководитель:", "________________"), ("", "Задание принял:", "________________")]
    for i, row in enumerate(sign.rows):
        for j, cell in enumerate(row.cells):
            base.set_cell_width(cell, [5.5, 4.5, 6.5][j])
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(sign_values[i][j])
            base.set_run_font(r, size=12.5)
    base.remove_table_borders(sign)
    doc.add_page_break()


def _style_table_cell(cell, text: str, width: float, *, bold=False, fill=None, size=11.5, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    base.set_cell_width(cell, width)
    base.set_cell_margins(cell, 45, 70, 45, 70)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        base.set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text or " ")
    base.set_run_font(r, size=size, bold=bold)


def add_schedule_page(doc: Document) -> None:
    base.add_plain_paragraph(doc, "8. Консультанты по отдельным разделам ДП:", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, bold=True, line=1.0, space_after=8)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [3.7, 4.7, 4.0, 4.1]
    headers = ["Раздел", "ФИО руководителя", "Задание выдал", "Задание принял"]
    for j, value in enumerate(headers):
        _style_table_cell(table.rows[0].cells[j], value, widths[j], bold=True, fill=base.PALE_BLUE, size=11)
    base.set_repeat_table_header(table.rows[0])
    rows = [
        ["Введение", "Азамова С. Ф.", "", ""],
        ["Глава 1", "Азамова С. Ф.", "", ""],
        ["Глава 2", "Азамова С. Ф.", "", ""],
        ["Глава 3", "Азамова С. Ф.", "", ""],
        ["Глава 4. БЖД", "", "", ""],
        ["Заключение", "Азамова С. Ф.", "", ""],
    ]
    for values in rows:
        row = table.add_row()
        base.prevent_row_split(row)
        for j, value in enumerate(values):
            _style_table_cell(row.cells[j], value, widths[j], size=11, align=WD_ALIGN_PARAGRAPH.LEFT if j < 2 else WD_ALIGN_PARAGRAPH.CENTER)

    base.add_plain_paragraph(doc, "9. График выполнения работы:", align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, bold=True, line=1.0, space_before=16, space_after=8)
    schedule = doc.add_table(rows=1, cols=4)
    schedule.style = "Table Grid"
    schedule.alignment = WD_TABLE_ALIGNMENT.CENTER
    schedule.autofit = False
    widths = [1.0, 8.2, 3.5, 3.8]
    headers = ["№", "Наименование раздела", "Срок", "Подпись"]
    for j, value in enumerate(headers):
        _style_table_cell(schedule.rows[0].cells[j], value, widths[j], bold=True, fill=base.PALE_BLUE, size=11)
    tasks = [
        "Глава 1. Анализ предметной области и постановка задачи",
        "Глава 2. Проектирование комплекса UFI Phone",
        "Глава 3. Реализация и тестирование",
        "Глава 4. Безопасность жизнедеятельности",
        "Заключение и оформление",
        "Подготовка презентации и предварительная защита",
    ]
    for index, task in enumerate(tasks, 1):
        row = schedule.add_row()
        base.prevent_row_split(row)
        for j, value in enumerate([str(index), task, "", ""]):
            _style_table_cell(row.cells[j], value, widths[j], size=10.8, align=WD_ALIGN_PARAGRAPH.LEFT if j == 1 else WD_ALIGN_PARAGRAPH.CENTER)

    signatures = doc.add_table(rows=2, cols=3)
    signatures.alignment = WD_TABLE_ALIGNMENT.CENTER
    signatures.autofit = False
    sign_rows = [
        ("Выпускник:", "________________", "«_____» __________ 2026 г."),
        ("Руководитель:", "________________", "«_____» __________ 2026 г."),
    ]
    for i, row in enumerate(signatures.rows):
        for j, cell in enumerate(row.cells):
            base.set_cell_width(cell, [2.8, 3.3, 10.4][j])
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(15 if i == 0 else 5)
            r = p.add_run(sign_rows[i][j])
            base.set_run_font(r, size=11.5)
    base.remove_table_borders(signatures)
    doc.add_page_break()


def add_manual_toc(doc: Document) -> None:
    for level, title, page in TOC_ENTRIES:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65 if level else 0)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(f"{title}\t{page}")
        base.set_run_font(r, size=10.5, bold=level == 0)


def add_front_matter(doc: Document) -> None:
    base.add_centered(doc, "СОДЕРЖАНИЕ", bold=True, space_after=10)
    add_manual_toc(doc)
    doc.add_page_break()

    base.add_centered(doc, "АННОТАЦИЯ", bold=True, space_after=10)
    annotations = [
        "В дипломном проекте разработан программно-аппаратный комплекс UFI Phone, превращающий совместимый LTE USB-модем UFI003 в локальный шлюз обычных сотовых вызовов и SMS. Привилегированные службы на Android 4.4 управляют радиорежимом, телефонией, базой сообщений и аудиоканалом. Android-, Linux- и Windows-клиенты подключаются по частной сети модема с аутентификацией challenge-response на основе HMAC-SHA-256. Реализованы входящие и исходящие вызовы, двусторонний звук, диалоги SMS, локальная история, автоматический запуск и восстановление режима LTE/3G fallback. Установка разрешается только при точном совпадении проверенного аппаратного профиля. Результат представляет работающий прототип версии 0.5.0, испытанный с оператором Ucell и подготовленный к расширению через воспроизводимую матрицу приёмки.",
        "Diplom loyihasida mos LTE USB UFI003 modemini oddiy uyali qo‘ng‘iroqlar va SMS uchun mahalliy shlyuzga aylantiruvchi UFI Phone dasturiy-apparat majmuasi ishlab chiqildi. Android 4.4 modemidagi imtiyozli xizmatlar radio rejimi, telefoniya, xabarlar bazasi va audio kanalini boshqaradi. Android, Linux va Windows mijozlari modemning yopiq tarmog‘i orqali HMAC-SHA-256 asosidagi challenge-response autentifikatsiyasi bilan ulanadi. Kiruvchi va chiquvchi qo‘ng‘iroqlar, ikki tomonlama ovoz, SMS dialoglari, mahalliy tarix, avtomatik ishga tushish hamda LTE/3G fallback rejimini tiklash amalga oshirilgan. O‘rnatish faqat tekshirilgan apparat profiliga to‘liq mos kelganda ruxsat etiladi. Natija — Ucell operatorida sinovdan o‘tgan va yangi profil hamda operatorlarni qabul qilish matritsasi orqali kengaytiriladigan 0.5.0 versiyadagi ishlaydigan prototip.",
        "This diploma project presents UFI Phone, a hardware–software system that turns a compatible UFI003 LTE USB modem into a private gateway for ordinary cellular calls and SMS. Privileged Android 4.4 services control radio mode, telephony, the message store, and bidirectional audio. Android, Linux, and Windows clients connect through the modem’s private LAN using an HMAC-SHA-256 challenge-response protocol. The implementation supports incoming and outgoing calls, two-way audio, SMS conversations, local history, automatic startup, and recovery of LTE/3G fallback mode. Installation is permitted only after an exact tested hardware-profile match. The result is a working version 0.5.0 prototype validated with Ucell and structured for controlled expansion through a reproducible acceptance matrix.",
    ]
    for index, text in enumerate(annotations):
        base.add_plain_paragraph(doc, text, size=11.5, line=1.15, space_before=0 if index == 0 else 6)
    doc.add_page_break()


def _box(draw: ImageDraw.ImageDraw, xy, text: str, font, *, fill="#EAF2F8", outline="#315F7D", text_fill="#111111") -> None:
    base.rounded(draw, xy, text, fill=fill, outline=outline, font=font, text_fill=text_fill)


def make_figures() -> dict[str, Path]:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    figures: dict[str, Path] = {}
    title_font = ImageFont.truetype(base.FONT_BOLD, 42)
    bold = ImageFont.truetype(base.FONT_BOLD, 31)
    regular = ImageFont.truetype(base.FONT_REGULAR, 28)
    small = ImageFont.truetype(base.FONT_REGULAR, 24)

    img, d = base.new_canvas("Контекст UFI Phone")
    clients = [
        (90, 220, 510, 380, "Android-планшет\nUFI Phone"),
        (90, 455, 510, 615, "Linux / Windows\ndesktop client"),
        (90, 690, 510, 850, "CLI / installer\nADB и USB"),
    ]
    for box in clients:
        _box(d, box[:4], box[4], bold, fill="#F8FBFD")
        base.arrow(d, (box[2], (box[1] + box[3]) // 2), (690, 535))
    _box(d, (690, 330, 1160, 740), "UFI003\nAndroid 4.4\n\nNetwork Guard\nVoice/SMS Gateway\nSIM и radio", bold, fill="#DCEAF3")
    base.arrow(d, (1160, 535), (1325, 535))
    _box(d, (1325, 355, 1720, 715), "Сеть оператора\n\nLTE — данные\nWCDMA/HSPA — речь\nSMS — carrier", bold, fill="#FFF4D6", outline="#B48216")
    d.text((120, 905), "Частная LAN 192.168.100.0/24 · без обязательного облачного relay", font=regular, fill="#244A66")
    p = WORK_DIR / "figure_context.png"
    img.save(p, quality=95)
    figures["context"] = p

    img, d = base.new_canvas("Радиопереход при вызове")
    stages = [
        (90, 350, 420, 670, "Ожидание\nLTE data\nmode 9", "#DFF2E5", "#2D7A47"),
        (520, 350, 850, 670, "RINGING / DIALING\nсеть инициирует\nfallback", "#EAF2F8", "#315F7D"),
        (950, 350, 1280, 670, "ACTIVE\nWCDMA / HSPA\ncircuit-switched", "#FFF4D6", "#B48216"),
        (1380, 350, 1710, 670, "IDLE\nвозврат LTE\nmobile data", "#DFF2E5", "#2D7A47"),
    ]
    for index, (x0, y0, x1, y1, text, fill, outline) in enumerate(stages):
        _box(d, (x0, y0, x1, y1), text, bold, fill=fill, outline=outline)
        if index < len(stages) - 1:
            base.arrow(d, (x1, 510), (stages[index + 1][0], 510))
    d.text((210, 790), "Network Guard исправляет LTE-only (mode 11) → автоматический LTE/GSM/WCDMA (mode 9)", font=regular, fill="#244A66")
    p = WORK_DIR / "figure_radio.png"
    img.save(p, quality=95)
    figures["radio"] = p

    img, d = base.new_canvas("Архитектура комплекса")
    layers = [
        (100, 185, 1700, 340, "Клиенты: Android 8+ · Linux · Windows\nCalls · Keypad · Messages · local history · notifications", "#F8FBFD"),
        (100, 390, 1700, 545, "Локальный протокол: control 8765 · downlink 8766 · uplink 8767\nHMAC-SHA-256 · nonce · private /24 · one-audio-client", "#EAF2F8"),
        (100, 595, 1700, 750, "Службы модема: Voice Gateway · SMS Controller · Network Guard\nboot recovery · telephony APIs · audio bridge", "#DCEAF3"),
        (100, 800, 1700, 955, "Аппаратный и сетевой слой: UFI003 · Qualcomm MSM8916 · SIM · carrier radio", "#FFF4D6"),
    ]
    for box in layers:
        _box(d, box[:4], box[4], bold, fill=box[4 + 1] if False else box[5])
    for y in (360, 565, 770):
        base.arrow(d, (900, y - 10), (900, y + 20))
    p = WORK_DIR / "figure_architecture.png"
    img.save(p, quality=95)
    figures["architecture"] = p

    img, d = base.new_canvas("Граница доверия и каналы")
    d.rounded_rectangle((65, 155, 1735, 970), radius=28, fill="#F8FBFD", outline="#66717A", width=4)
    d.text((95, 175), "Доверенная частная сеть модема 192.168.100.0/24", font=bold, fill="#244A66")
    _box(d, (130, 330, 555, 770), "Клиент\n\nPairing token\nв private storage\n\nAndroid / Desktop", bold, fill="#EAF2F8")
    _box(d, (755, 285, 1190, 815), "Gateway\n192.168.100.1\n\n8765 control\n8766 downlink\n8767 uplink\n\nMAX_CLIENTS=12", bold, fill="#DCEAF3")
    _box(d, (1390, 350, 1670, 750), "SIM / radio\n\nSMS DB\nTelephony\nAudio route", bold, fill="#FFF4D6", outline="#B48216")
    for y, label in [(410, "HELLO 2 + nonce"), (550, "AUTH HMAC proof"), (690, "bound command/stream")]:
        base.arrow(d, (555, y), (755, y))
        d.text((565, y - 55), label, font=small, fill="#244A66")
    base.arrow(d, (1190, 550), (1390, 550))
    d.text((110, 890), "ADB shell-only receivers configure secrets; clients outside /24 are rejected before command handling", font=small, fill="#244A66")
    p = WORK_DIR / "figure_trust.png"
    img.save(p, quality=95)
    figures["trust"] = p

    img, d = base.new_canvas("Профильная установка")
    nodes = [
        (80, 230, 380, 440, "ufi_setup.py\ndoctor\nread-only"),
        (500, 230, 800, 440, "Сбор whitelist:\nproduct, SDK,\nbaseband, USB, LAN"),
        (920, 230, 1220, 440, "Exact match\nhardware-profiles.json"),
        (1340, 150, 1710, 360, "MATCH\nсборка и проверка\nсертификата"),
        (1340, 570, 1710, 780, "NO MATCH\nотказ до установки\nprivileged APK"),
    ]
    colors = ["#F8FBFD", "#EAF2F8", "#DCEAF3", "#DFF2E5", "#FDE7E7"]
    outlines = ["#315F7D", "#315F7D", "#315F7D", "#2D7A47", "#A33A3A"]
    for node, fill, outline in zip(nodes, colors, outlines):
        _box(d, node[:4], node[4], bold, fill=fill, outline=outline)
    base.arrow(d, (380, 335), (500, 335))
    base.arrow(d, (800, 335), (920, 335))
    base.arrow(d, (1220, 300), (1340, 255))
    base.arrow(d, (1220, 390), (1340, 675), color="#A33A3A")
    _box(d, (500, 690, 1220, 915), "После MATCH: Network Guard + Voice Gateway → modem\nUFI Phone → tablet · token → только через ADB\nTelegram остаётся выключенным", regular, fill="#F8FBFD")
    p = WORK_DIR / "figure_install.png"
    img.save(p, quality=95)
    figures["install"] = p

    img, d = base.new_canvas("Входящий вызов")
    lanes = [(140, "Оператор"), (560, "UFI003 Gateway"), (1020, "Monitor Service"), (1450, "Пользователь")]
    for x, label in lanes:
        d.text((x - 70, 170), label, font=bold, fill="#244A66")
        d.line((x, 245, x, 950), fill="#9FB6C5", width=4)
    events = [
        (140, 560, 320, "RINGING + номер"),
        (560, 1020, 430, "STATUS: RINGING"),
        (1020, 1450, 540, "full-screen alert"),
        (1450, 1020, 650, "Answer"),
        (1020, 560, 760, "ANSWER + HMAC"),
        (560, 140, 870, "telephony answer"),
    ]
    for x0, x1, y, label in events:
        base.arrow(d, (x0, y), (x1, y))
        d.text((min(x0, x1) + 25, y - 44), label, font=small, fill="#111111")
    p = WORK_DIR / "figure_call_sequence.png"
    img.save(p, quality=95)
    figures["call_sequence"] = p

    img, d = base.new_canvas("Двусторонний аудиотракт")
    _box(d, (90, 300, 460, 760), "Сотовая сеть\nи modem audio\n\nVOICE_DOWNLINK\nFIFO uplink", bold, fill="#FFF4D6", outline="#B48216")
    _box(d, (715, 245, 1085, 815), "Voice Gateway\n\n8766: 8 kHz mono\n8767: 48 kHz mono\n\nS16LE · busy lock", bold, fill="#DCEAF3")
    _box(d, (1340, 300, 1710, 760), "Клиент\n\nAudioTrack speaker\nAudioRecord mic\nEcho canceller", bold, fill="#EAF2F8")
    base.arrow(d, (460, 420), (715, 420))
    base.arrow(d, (1085, 420), (1340, 420))
    d.text((505, 365), "downlink", font=small, fill="#244A66")
    base.arrow(d, (1340, 650), (1085, 650))
    base.arrow(d, (715, 650), (460, 650))
    d.text((1110, 595), "uplink", font=small, fill="#244A66")
    base.draw_wrapped(
        d,
        (250, 850, 1550, 985),
        "Аудиосессия разрешена только в состоянии ACTIVE; завершение вызова закрывает оба потока",
        small,
        fill="#244A66",
        align="center",
        spacing=4,
    )
    p = WORK_DIR / "figure_audio.png"
    img.save(p, quality=95)
    figures["audio"] = p

    img, d = base.new_canvas("Поток SMS")
    nodes = [
        (90, 350, 400, 690, "Carrier SMS\n\nSMS_RECEIVED"),
        (510, 350, 820, 690, "SmsReceiver\nи SmsDb\n\nнормализация"),
        (930, 350, 1240, 690, "Gateway\nSMS_LIST\nSMS_SEND\nSMS_READ"),
        (1350, 350, 1710, 690, "UFI Phone\nдиалоги\nуведомления\nдо 250 записей"),
    ]
    for i, node in enumerate(nodes):
        _box(d, node[:4], node[4], bold, fill=["#FFF4D6", "#DCEAF3", "#EAF2F8", "#F8FBFD"][i], outline="#315F7D")
        if i < len(nodes) - 1:
            base.arrow(d, (node[2], 520), (nodes[i + 1][0], 520))
    d.text((295, 820), "Основной путь не использует Telegram; legacy forwarder включается отдельно и явно", font=regular, fill="#244A66")
    p = WORK_DIR / "figure_sms.png"
    img.save(p, quality=95)
    figures["sms"] = p

    calls = Image.open(ASSET_DIR / "ufi-app-calls.png").convert("RGB")
    messages = Image.open(ASSET_DIR / "ufi-app-messages.png").convert("RGB")
    canvas, d = base.new_canvas("Интерфейс UFI Phone", size=(1800, 1120))
    calls = ImageOps.contain(calls, (760, 770), Image.Resampling.LANCZOS)
    messages = ImageOps.contain(messages, (760, 770), Image.Resampling.LANCZOS)
    canvas.paste(ImageOps.expand(calls, border=4, fill="#9FB6C5"), (95, 220))
    canvas.paste(ImageOps.expand(messages, border=4, fill="#9FB6C5"), (945, 220))
    d.text((95, 175), "Вызовы и готовность сети", font=regular, fill="#244A66")
    d.text((945, 175), "Диалоги и отправка SMS", font=regular, fill="#244A66")
    p = WORK_DIR / "figure_ui.png"
    canvas.save(p, quality=95)
    figures["ui"] = p

    img, d = base.new_canvas("Поставка компонентов")
    columns = [
        (90, 235, 555, 840, "Модем\n\nNetwork Guard APK\nVoice Gateway APK\n\nplatform-signed\nprofile-gated ADB\nне публикуется в store", "#FFF4D6"),
        (665, 235, 1130, 840, "Android companion\n\nAPK / AAB\nAndroid 8+\nAPI target 36\n\nGoogle Play materials\nand privacy policy", "#EAF2F8"),
        (1240, 235, 1710, 840, "Desktop\n\nPython source\nWindows executable\nLinux portable/deb\n\nlocal config\nand audio backend", "#DCEAF3"),
    ]
    for box in columns:
        _box(d, box[:4], box[4], bold, fill=box[5], outline="#315F7D")
    p = WORK_DIR / "figure_deployment.png"
    img.save(p, quality=95)
    figures["deployment"] = p

    img, d = base.new_canvas("Проверка версии 0.5.0")
    _box(d, (100, 235, 780, 850), "Автоматические проверки\n\n12 / 12 unittest passed\n1,022 s\n\nmanifest consistency\nmodel and history\nprofile rejection\nSMS durability\nvoice protocol", bold, fill="#DFF2E5", outline="#2D7A47")
    _box(d, (1020, 235, 1700, 850), "Repository guard\n\nversion 0.5.0 aligned\nsecret patterns checked\nfile sizes checked\nGitHub Actions pinned\n\nRESULT: PASSED", bold, fill="#EAF2F8", outline="#315F7D")
    base.arrow(d, (780, 540), (1020, 540))
    p = WORK_DIR / "figure_quality.png"
    img.save(p, quality=95)
    figures["quality"] = p
    return figures


def build_body(doc: Document, figures: dict[str, Path]) -> None:
    blocks = re.split(r"\n\s*\n", CONTENT_PATH.read_text(encoding="utf-8").strip())
    for block in blocks:
        block = block.strip().replace("`", "")
        if not block or block.startswith("<!--"):
            continue
        if block == "[[PAGEBREAK]]":
            doc.add_page_break()
            continue
        directive = re.fullmatch(r"\[\[(TABLE|FIGURE):([^|\]]+)\|([^\]]+)\]\]", block)
        if directive:
            kind, number, key = directive.groups()
            if kind == "TABLE":
                base.add_table(doc, number.strip(), key.strip())
            else:
                key = key.strip()
                base.add_figure(doc, number.strip(), figures[key], FIGURE_CAPTIONS[key])
            continue
        if block.startswith("# "):
            if doc.paragraphs and doc.paragraphs[-1].text.strip():
                doc.add_page_break()
            p = doc.add_paragraph(style="Heading 1")
            base.parse_inline(p, block[2:].strip())
            continue
        if block.startswith("## "):
            p = doc.add_paragraph(style="Heading 2")
            base.parse_inline(p, block[3:].strip())
            continue
        if block.startswith("### "):
            p = doc.add_paragraph(style="Heading 3")
            base.parse_inline(p, block[4:].strip())
            continue
        if block.startswith("- "):
            for line in block.splitlines():
                base.add_list_item(doc, line[2:].strip())
            continue
        base.add_body_paragraph(doc, " ".join(line.strip() for line in block.splitlines()))


SOURCES = [
    ("Закон Республики Узбекистан «О телекоммуникациях» от 20.08.1999 № 822-I.", "https://lex.uz/docs/33152"),
    ("Закон Республики Узбекистан «О персональных данных» от 02.07.2019 № ЗРУ-547.", "https://lex.uz/ru/docs/4396419"),
    ("Закон Республики Узбекистан «Об охране труда» в редакции Закона от 22.09.2016 № ЗРУ-410.", "https://lex.uz/ru/docs/3031427"),
    ("Закон Республики Узбекистан «О пожарной безопасности» от 30.09.2009 № ЗРУ-226.", "https://lex.uz/ru/docs/1521661"),
    ("Указ Президента Республики Узбекистан от 05.10.2020 № УП-6079 «Цифровой Узбекистан — 2030».", "https://lex.uz/ru/docs/5031048"),
    ("Национальный комитет Республики Узбекистан по статистике. Социально-экономическое положение Республики Узбекистан за январь–март 2025 года.", "https://stat.uz/files/538/2025-Choraklik-natijalar-january--march-ang/3940/Report-for-January-March-2025.pdf"),
    ("Комитет по развитию конкуренции и защите прав потребителей. Результаты анализа рынка услуг мобильной связи.", "https://raqobat.gov.uz/ru/rezultaty-analiza-rynka-uslug-mobilnoj-svyazi/"),
    ("Ucell. VoLTE/ViLTE — новые возможности технологии.", "https://ucell.uz/en/company_news/ucell_volte_vilte-_texnologiyalarining_yangi_ufqlari"),
    ("Mobiuz. Voice over LTE (VoLTE).", "https://corp.mobi.uz/en/uslugi/volte/?VOICE=Y"),
    ("Uztelecom. Voice over LTE service.", "https://uztelecom.uz/en/for-individuals/mobile-communication/gsm/services/additional-services/volte/"),
    ("Beeline Uzbekistan. VoLTE requirements.", "https://b2b.beeline.uz/en/products/services/volte"),
    ("Humans. Telecom network information.", "https://qr.humans.uz/ru/telecom"),
    ("Perfectum. CDMA network information.", "https://perfectum.uz/uz/cdma"),
    ("Android Developers. TelephonyManager API reference.", "https://developer.android.com/reference/android/telephony/TelephonyManager"),
    ("Android Developers. PhoneStateListener API reference.", "https://developer.android.com/reference/android/telephony/PhoneStateListener"),
    ("Android Developers. SmsManager API reference.", "https://developer.android.com/reference/android/telephony/SmsManager"),
    ("Android Developers. AudioRecord API reference.", "https://developer.android.com/reference/android/media/AudioRecord"),
    ("Android Developers. AudioTrack API reference.", "https://developer.android.com/reference/android/media/AudioTrack"),
    ("Android Developers. Foreground services overview.", "https://developer.android.com/develop/background-work/services/fgs"),
    ("Android Open Source Project. Application sandbox and platform security model.", "https://source.android.com/docs/security/app-sandbox"),
    ("Krawczyk H., Bellare M., Canetti R. HMAC: Keyed-Hashing for Message Authentication. RFC 2104, 1997.", "https://www.rfc-editor.org/rfc/rfc2104"),
    ("NIST. FIPS PUB 180-4: Secure Hash Standard (SHS).", "https://csrc.nist.gov/pubs/fips/180-4/upd1/final"),
    ("OWASP. Authentication Cheat Sheet.", "https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html"),
    ("3GPP TS 23.272. Circuit Switched Fallback in Evolved Packet System; Stage 2.", "https://www.3gpp.org/dynareport/23272.htm"),
    ("3GPP TS 27.007. AT command set for User Equipment.", "https://www.3gpp.org/dynareport/27007.htm"),
    ("3GPP TS 27.005. Use of Data Terminal Equipment–Data Circuit terminating Equipment interface for SMS and CBS.", "https://www.3gpp.org/dynareport/27005.htm"),
    ("ISO/IEC 25010:2023. Systems and software Quality Requirements and Evaluation — Product quality model.", "https://www.iso.org/standard/78176.html"),
    ("Sommerville I. Software Engineering. 10th ed. Pearson, 2015. 816 p.", None),
    ("Pressman R. S., Maxim B. R. Software Engineering: A Practitioner’s Approach. 9th ed. McGraw-Hill, 2020. 704 p.", None),
    ("Bass L., Clements P., Kazman R. Software Architecture in Practice. 4th ed. Addison-Wesley, 2021. 464 p.", None),
    ("UFI Phone. README и описание возможностей проекта. Локальная документация репозитория, версия 0.5.0, 2026.", None),
    ("UFI Phone. Product vision. Локальная архитектурная документация, редакция 2026.", None),
    ("UFI Phone. Adding support for another UFI modem. Руководство по аппаратным профилям, 2026.", None),
    ("UFI Phone. Uzbekistan market and operator compatibility. Исследование от 24.09.2026.", None),
    ("UFI Phone. Privacy Policy and Data Safety. Материалы Google Play, редакция 2026.", None),
    ("Python Software Foundation. Python 3 documentation.", "https://docs.python.org/3/"),
    ("PyUSB documentation.", "https://pyusb.github.io/pyusb/"),
    ("Gradle. Build Tool documentation.", "https://docs.gradle.org/"),
    ("GitHub Docs. GitHub Actions documentation.", "https://docs.github.com/actions"),
]


def add_bibliography(doc: Document) -> None:
    doc.add_page_break()
    p = doc.add_paragraph(style="Heading 1")
    base.parse_inline(p, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    for index, (text, url) in enumerate(SOURCES, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{index}. {text}")
        base.set_run_font(r, size=12)
        if url:
            r = p.add_run(" URL: ")
            base.set_run_font(r, size=12)
            base.add_hyperlink(p, url, url)
            r = p.add_run(" (дата обращения: 24.09.2026).")
            base.set_run_font(r, size=12)


def audit(doc: Document) -> None:
    table_text = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    text = "\n".join([*(p.text for p in doc.paragraphs), *table_text])
    required = [
        "Куватов Руслан Бахтиярович",
        "60610500",
        "ГЛАВА I.",
        "ГЛАВА II.",
        "ГЛАВА III.",
        "ГЛАВА IV.",
        "ЗАКЛЮЧЕНИЕ",
        "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
        "UFI003",
        "HMAC-SHA-256",
    ]
    missing = [value for value in required if value not in text]
    if missing:
        raise RuntimeError(f"Missing required content: {missing}")
    if len(doc.inline_shapes) < 10:
        raise RuntimeError(f"Expected at least 10 figures, got {len(doc.inline_shapes)}")
    if len(doc.tables) < 12:
        raise RuntimeError(f"Expected at least 12 tables, got {len(doc.tables)}")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    figures = make_figures()
    doc = Document()
    base.configure_document(doc)
    doc.core_properties.title = "Программно-аппаратный комплекс UFI Phone"
    doc.core_properties.subject = "Дипломный проект"
    doc.core_properties.author = "Куватов Руслан Бахтиярович"
    doc.core_properties.keywords = "UFI Phone, UFI003, LTE USB modem, SMS, cellular calls, Android"
    doc.core_properties.comments = ""
    doc.core_properties.last_modified_by = "Куватов Руслан Бахтиярович"
    doc.core_properties.created = datetime(2026, 9, 25)
    doc.core_properties.modified = datetime(2026, 9, 26)
    add_title_page(doc)
    add_assignment_page(doc)
    add_schedule_page(doc)
    add_front_matter(doc)
    build_body(doc, figures)
    add_bibliography(doc)
    audit(doc)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
